"""Photos, videos and documents posted to a diary.

The apps encrypt every media file with the space key into one binary blob (FILE_ATTACHMENT_SPEC):

    [nonce length: uint32 little-endian][nonce, 12 bytes][ciphertext || 16-byte GCM tag]

Images are JPEG, videos MP4. A document's name, MIME type and size are not on the server in
plaintext either: they travel in ``encrypted_metadata.file_info``, base64 of the same blob
layout around a small JSON object.

Everything here runs after download, on the machine running the agent. Nothing is sent
anywhere except to the backend the operator chose.
"""

from __future__ import annotations

import base64
import json
import logging
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from cryptography.hazmat.primitives.ciphers.aead import AESGCM

log = logging.getLogger(__name__)

#: How much of a document's text is handed to a model. A diary partner reading a document
#: for a grammar check needs the words, not a 400-page manual.
MAX_TEXT_CHARS = 60_000

TEXT_MIMES = {"text/plain", "text/markdown", "text/csv", "application/json"}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".tsv", ".log"}


@dataclass
class Attachment:
    """A media file on a diary entry. ``path`` is set once it has been downloaded."""

    kind: str  # "image" | "video" | "file"
    message_uuid: str
    chat_space_id: str
    name: str = ""
    mime: str = ""
    size: int = 0
    path: Path | None = None
    #: The message's ``encrypted_metadata`` (nonce, and for documents the sealed file_info).
    metadata: dict = field(default_factory=dict, repr=False)

    @property
    def label(self) -> str:
        return {"image": "photo", "video": "video"}.get(self.kind) or (self.name or "file")

    def as_json(self) -> dict:
        return {"kind": self.kind, "name": self.name, "mime": self.mime, "size": self.size,
                "path": str(self.path) if self.path else None}


def decrypt_blob(blob: bytes, key: bytes) -> bytes:
    """Open the apps' media blob. Raises ValueError on a malformed or undecryptable blob."""
    if len(blob) < 4 + 12 + 16:
        raise ValueError("media blob is too short")
    n = int.from_bytes(blob[:4], "little")
    if n <= 0 or n > 64:
        raise ValueError(f"implausible nonce length {n}")
    nonce, sealed = blob[4:4 + n], blob[4 + n:]
    try:
        return AESGCM(key).decrypt(nonce, sealed, None)
    except Exception as e:  # noqa: BLE001 - InvalidTag and friends
        raise ValueError(f"could not decrypt media: {e}") from e


def parse_metadata(value) -> dict:  # noqa: ANN001
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value:
        try:
            return json.loads(value)
        except ValueError:
            return {}
    return {}


def file_info(metadata: dict, key: bytes) -> dict:
    """The document's ``{"name", "mime", "size"}``, or {} when absent or unreadable."""
    raw = metadata.get("file_info")
    if not raw:
        return {}
    try:
        return json.loads(decrypt_blob(base64.b64decode(raw), key).decode("utf-8"))
    except Exception as e:  # noqa: BLE001
        log.warning("could not read file_info: %s", e)
        return {}


def attachment_from_message(message_type: str, message_uuid: str, chat_space_id: str,
                            metadata: dict) -> Attachment | None:
    if message_type not in ("image", "video", "file") or not message_uuid:
        return None
    mime = {"image": "image/jpeg", "video": "video/mp4"}.get(message_type, "")
    return Attachment(kind=message_type, message_uuid=message_uuid,
                      chat_space_id=chat_space_id, mime=mime, metadata=metadata)


def sniff_image(data: bytes) -> tuple[str, str] | None:
    """(mime, suffix) from an image's first bytes. The apps send JPEG, but a robot's camera
    frames are PNG, and a model rejects an image declared as the wrong type."""
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg", ".jpg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png", ".png"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif", ".gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", ".webp"
    if data[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1"):
        return "image/heic", ".heic"
    return None


def suffix_for(att: Attachment) -> str:
    if att.kind == "image":
        return ".jpg"
    if att.kind == "video":
        return ".mp4"
    s = Path(att.name).suffix.lower()
    return s if s and len(s) <= 6 else ".bin"


# -- turning a file into something a model can take -------------------------------------------

def video_frames(path: Path, count: int = 4, width: int = 768) -> list[Path]:
    """Evenly spaced stills from a video, via ffmpeg. [] when ffmpeg is not installed.

    Models read images, not video; a handful of frames is how a video is shown to them.
    """
    ffmpeg, ffprobe = shutil.which("ffmpeg"), shutil.which("ffprobe")
    if not ffmpeg or not ffprobe:
        return []
    try:
        duration = float(subprocess.run(
            [ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
             str(path)], capture_output=True, text=True, timeout=30).stdout.strip() or 0)
    except (ValueError, OSError, subprocess.TimeoutExpired):
        duration = 0.0
    frames = []
    for i in range(count):
        at = duration * (i + 0.5) / count if duration > 0 else 0
        out = path.with_name(f"{path.stem}_frame{i + 1}.jpg")
        if not out.exists():
            subprocess.run([ffmpeg, "-v", "error", "-y", "-ss", f"{at:.2f}", "-i", str(path),
                            "-frames:v", "1", "-vf", f"scale={width}:-2", str(out)],
                           capture_output=True, timeout=60)
        if out.exists():
            frames.append(out)
        if duration <= 0:
            break
    return frames


def document_text(path: Path, mime: str = "") -> str | None:
    """The text of a document, or None when this format cannot be read here.

    Plain text needs nothing. Word, Excel, PowerPoint and PDF use optional libraries
    (``pip install 'pyunto-agent[docs]'``); without them the model is told the file exists but
    not what is in it, which is honest rather than silently empty.
    """
    suffix = path.suffix.lower()
    try:
        if mime in TEXT_MIMES or suffix in TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]
        if suffix == ".docx":
            import docx  # noqa: PLC0415 - optional

            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)[:MAX_TEXT_CHARS]
        if suffix == ".pptx":
            from pptx import Presentation  # noqa: PLC0415 - optional

            lines = []
            for n, slide in enumerate(Presentation(str(path)).slides, 1):
                lines.append(f"--- slide {n} ---")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text)
            return "\n".join(lines)[:MAX_TEXT_CHARS]
        if suffix == ".xlsx":
            from openpyxl import load_workbook  # noqa: PLC0415 - optional

            lines = []
            for ws in load_workbook(str(path), read_only=True, data_only=True).worksheets:
                lines.append(f"--- sheet {ws.title} ---")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        lines.append("\t".join("" if c is None else str(c) for c in row))
            return "\n".join(lines)[:MAX_TEXT_CHARS]
        if suffix == ".pdf":
            from pypdf import PdfReader  # noqa: PLC0415 - optional

            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)[:MAX_TEXT_CHARS]
    except ImportError:
        log.info("reading %s needs the docs extra: pip install 'pyunto-agent[docs]'", suffix)
        return None
    except Exception as e:  # noqa: BLE001 - a broken document must not stop the reply
        log.warning("could not read %s: %s", path.name, e)
        return None
    return None
